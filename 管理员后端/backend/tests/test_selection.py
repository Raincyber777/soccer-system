# -*- coding: utf-8 -*-
"""Phase 4 遴选域测试：评分录入/修改、乐观锁、汇总排名、导出、结果生成。"""
from __future__ import annotations

import csv
import io
import zipfile
from decimal import Decimal

from django.core.cache import cache
from django.test import Client, TestCase

from apps.accounts.models import User
from apps.core.models import Region
from apps.oplogs.models import OperationLog
from apps.players.models import Player
from apps.selection.models import Event, SelectionResult, SelectionScore
from apps.teams.models import Team


def build_score(event, player, judge, stage, total, **kwargs):
    payload = dict(
        event=event, player=player, team_id=player.team_id, category=player.category,
        stage=stage, judge=judge, judge_name=judge.real_name or judge.username,
        physical_score=Decimal("16.0"), technical_score=Decimal("16.0"),
        tactical_score=Decimal("16.0"), mental_score=Decimal("16.0"),
        potential_score=Decimal("16.0"), total_score=Decimal(str(total)), version=1,
    )
    payload.update(kwargs)
    return SelectionScore.objects.create(**payload)


class SelectionFixture(TestCase):
    def setUp(self) -> None:
        cache.clear()
        self.region = Region.objects.create(name="成都市", level=2, code="510100", sort=1)
        self.team = Team.objects.create(name="成都七中", region=self.region, category=2)
        self.other_team = Team.objects.create(name="绵阳南山中学", region=self.region, category=2)
        self.event = Event.objects.create(name="2026 中学组", year=2026, category=2)
        self.event_university = Event.objects.create(name="2026 高校组", year=2026, category=1)
        self.event_empty = Event.objects.create(name="2026 小学组", year=2026, category=3)

        self.admin = User.objects.create_user("admin", "admin123", role=9, status=1, real_name="系统管理员")
        self.judge_a = User.objects.create_user("judge_a", "admin123", role=9, status=1, real_name="评委甲")
        self.judge_b = User.objects.create_user("judge_b", "admin123", role=9, status=1, real_name="评委乙")

        self.players = []
        for index in range(6):
            team = self.team if index < 3 else self.other_team
            self.players.append(
                Player.objects.create(name=f"球员{index}", gender=1, team=team, category=2)
            )

        # 三名球员各两位评委评分，构造可预测排名
        build_score(self.event, self.players[0], self.judge_a, 2, "90.00")
        build_score(self.event, self.players[0], self.judge_b, 2, "92.00")
        build_score(self.event, self.players[1], self.judge_a, 2, "88.00")
        build_score(self.event, self.players[1], self.judge_b, 2, "88.00")
        build_score(self.event, self.players[2], self.judge_a, 2, "70.00")
        build_score(self.event, self.players[2], self.judge_b, 2, "72.00")
        build_score(self.event, self.players[0], self.judge_a, 1, "80.00")

        self.client = Client()
        self.client.post("/api/v1/auth/login", data={"username": "admin", "password": "admin123"}, content_type="application/json")


class ScoreCreateTests(SelectionFixture):
    def test_create_calculates_total(self) -> None:
        payload = {
            "event_id": self.event.id, "player_id": self.players[3].id, "stage": 1,
            "physical_score": 18.5, "technical_score": 17.0, "tactical_score": 16.5,
            "mental_score": 18.0, "potential_score": 17.5, "remark": "技术细腻",
        }
        body = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["total_score"], 87.5)
        self.assertEqual(body["data"]["judge_id"], self.admin.id)
        self.assertEqual(body["data"]["version"], 1)
        self.assertTrue(OperationLog.objects.filter(action="新增评分").exists())

    def test_client_total_ignored(self) -> None:
        payload = {
            "event_id": self.event.id, "player_id": self.players[3].id, "stage": 1,
            "physical_score": 10, "technical_score": 10, "tactical_score": 10,
            "mental_score": 10, "potential_score": 10, "total_score": 100,
        }
        body = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(body["data"]["total_score"], 50.0)

    def test_duplicate_returns_30003(self) -> None:
        payload = {
            "event_id": self.event.id, "player_id": self.players[3].id, "stage": 1,
            "physical_score": 15, "technical_score": 15, "tactical_score": 15,
            "mental_score": 15, "potential_score": 15,
        }
        first = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(first["code"], 0)
        # 同一评委 + 同一球员 + 同一阶段重复提交
        body = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(body["code"], 30003)

    def test_other_judge_on_same_player_allowed(self) -> None:
        payload = {
            "event_id": self.event.id, "player_id": self.players[3].id, "stage": 1,
            "physical_score": 15, "technical_score": 15, "tactical_score": 15,
            "mental_score": 15, "potential_score": 15,
        }
        self.assertEqual(
            self.client.post("/api/v1/selection-scores", data=payload, content_type="application/json").json()["code"],
            0,
        )
        other = Client()
        other.post("/api/v1/auth/login", data={"username": "judge_a", "password": "admin123"}, content_type="application/json")
        self.assertEqual(
            other.post("/api/v1/selection-scores", data=payload, content_type="application/json").json()["code"],
            0,
        )
        self.assertEqual(
            SelectionScore.objects.filter(player=self.players[3], stage=1).count(), 2
        )

    def test_category_mismatch_returns_40001(self) -> None:
        payload = {
            "event_id": self.event_university.id, "player_id": self.players[0].id, "stage": 1,
            "physical_score": 15, "technical_score": 15, "tactical_score": 15,
            "mental_score": 15, "potential_score": 15,
        }
        body = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(body["code"], 40001)
        self.assertIn("组别", body["msg"])

    def test_missing_field_returns_10001(self) -> None:
        body = self.client.post(
            "/api/v1/selection-scores",
            data={"event_id": self.event.id, "player_id": self.players[3].id, "stage": 1},
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 10001)

    def test_out_of_range_returns_10001(self) -> None:
        payload = {
            "event_id": self.event.id, "player_id": self.players[3].id, "stage": 1,
            "physical_score": 21, "technical_score": 15, "tactical_score": 15,
            "mental_score": 15, "potential_score": 15,
        }
        body = self.client.post(
            "/api/v1/selection-scores", data=payload, content_type="application/json"
        ).json()
        self.assertEqual(body["code"], 10001)


class ScoreUpdateTests(SelectionFixture):
    def test_update_success_increments_version(self) -> None:
        score = SelectionScore.objects.filter(player=self.players[0], judge=self.judge_a, stage=2).first()
        body = self.client.put(
            f"/api/v1/selection-scores/{score.id}",
            data={
                "physical_score": 19, "technical_score": 18, "tactical_score": 17,
                "mental_score": 19, "potential_score": 18, "remark": "复评微调", "version": score.version,
            },
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["version"], score.version + 1)
        self.assertEqual(body["data"]["total_score"], 91.0)
        log = OperationLog.objects.filter(action="修改评分").first()
        self.assertIn("before", log.params)
        self.assertEqual(log.params["before"]["total_score"], 90.0)

    def test_stale_version_returns_30008(self) -> None:
        score = SelectionScore.objects.filter(player=self.players[0], judge=self.judge_a, stage=2).first()
        body = self.client.put(
            f"/api/v1/selection-scores/{score.id}",
            data={
                "physical_score": 19, "technical_score": 18, "tactical_score": 17,
                "mental_score": 19, "potential_score": 18, "version": 99,
            },
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 30008)

    def test_missing_score_returns_30001(self) -> None:
        body = self.client.put(
            "/api/v1/selection-scores/999999",
            data={
                "physical_score": 19, "technical_score": 18, "tactical_score": 17,
                "mental_score": 19, "potential_score": 18, "version": 1,
            },
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 30001)


class SummaryTests(SelectionFixture):
    def test_ranking_and_overview(self) -> None:
        body = self.client.get(
            f"/api/v1/selection-scores/summary?event_id={self.event.id}&stage=2"
        ).json()
        self.assertEqual(body["code"], 0)
        data = body["data"]
        self.assertEqual(data["total"], 3)
        rows = data["list"]
        self.assertEqual([row["player_id"] for row in rows],
                         [self.players[0].id, self.players[1].id, self.players[2].id])
        self.assertEqual([row["rank"] for row in rows], [1, 2, 3])
        self.assertEqual(rows[0]["avg_score"], 91.0)
        self.assertEqual(rows[0]["judge_count"], 2)
        self.assertEqual(rows[0]["dimension_avg"]["physical_score"], 16.0)

        overview = data["overview"]
        self.assertEqual(overview["player_count"], 3)
        self.assertEqual(overview["judge_count"], 2)
        self.assertEqual(overview["max_score"], 92.0)
        self.assertEqual(overview["min_score"], 70.0)
        self.assertEqual(len(data["top10"]), 3)

    def test_pagination_keeps_global_rank(self) -> None:
        body = self.client.get(
            f"/api/v1/selection-scores/summary?event_id={self.event.id}&stage=2&pageNum=2&pageSize=1"
        ).json()
        self.assertEqual(body["data"]["list"][0]["rank"], 2)
        self.assertEqual(body["data"]["list"][0]["player_id"], self.players[1].id)

    def test_empty_summary_returns_zeros(self) -> None:
        body = self.client.get(
            f"/api/v1/selection-scores/summary?event_id={self.event_empty.id}&stage=1"
        ).json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["total"], 0)
        self.assertEqual(body["data"]["overview"]["player_count"], 0)
        self.assertEqual(body["data"]["top10"], [])

    def test_missing_event_returns_10001(self) -> None:
        body = self.client.get("/api/v1/selection-scores/summary?stage=1").json()
        self.assertEqual(body["code"], 10001)


class ExportTests(SelectionFixture):
    def test_export_xlsx(self) -> None:
        response = self.client.get(
            f"/api/v1/selection-scores/export?event_id={self.event.id}&stage=2&format=xlsx"
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn(".xlsx", response["Content-Disposition"])

        content = b"".join(response.streaming_content)
        self.assertTrue(zipfile.is_zipfile(io.BytesIO(content)))

        from openpyxl import load_workbook

        workbook = load_workbook(io.BytesIO(content), read_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "名次")
        self.assertEqual(len(rows), 4)  # 表头 + 3 条
        self.assertEqual(rows[1][0], 1)
        workbook.close()

        log = OperationLog.objects.filter(action="导出评分").first()
        self.assertIsNotNone(log)
        self.assertEqual(log.params["format"], "xlsx")
        self.assertEqual(log.params["row_count"], 3)

    def test_export_csv_with_bom(self) -> None:
        response = self.client.get(
            f"/api/v1/selection-scores/export?event_id={self.event.id}&stage=2&format=csv"
        )
        self.assertEqual(response.status_code, 200)
        content = b"".join(response.streaming_content)
        self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
        text = content.decode("utf-8-sig")
        rows = list(csv.reader(io.StringIO(text)))
        self.assertEqual(rows[0][0], "名次")
        self.assertEqual(len(rows), 4)

    def test_export_bad_format(self) -> None:
        body = self.client.get(
            f"/api/v1/selection-scores/export?event_id={self.event.id}&stage=2&format=pdf"
        ).json()
        self.assertEqual(body["code"], 10001)

    def test_export_without_event(self) -> None:
        body = self.client.get("/api/v1/selection-scores/export?stage=2").json()
        self.assertEqual(body["code"], 10001)


class GenerateTests(SelectionFixture):
    def test_generate_top5(self) -> None:
        build_score(self.event, self.players[3], self.judge_a, 2, "85.00")
        build_score(self.event, self.players[4], self.judge_a, 2, "60.00")
        build_score(self.event, self.players[5], self.judge_a, 2, "95.00")

        body = self.client.post(
            "/api/v1/selection-results/generate",
            data={"event_id": self.event.id, "stage": 2, "category": 2, "top_n": 5},
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["count"], 5)
        self.assertEqual(body["data"]["category_text"], "中学组")

        results = SelectionResult.objects.filter(event=self.event, stage=2, category=2, status=1).order_by("rank")
        self.assertEqual(results.count(), 5)
        self.assertEqual(results[0].player_id, self.players[5].id)
        self.assertEqual(results[0].rank, 1)
        self.assertEqual(float(results[0].total_score), 95.0)

        log = OperationLog.objects.filter(action="生成最佳球员结果").first()
        self.assertIsNotNone(log)
        self.assertEqual(log.params["count"], 5)

    def test_regenerate_invalidates_old(self) -> None:
        payload = {"event_id": self.event.id, "stage": 2, "category": 2, "top_n": 2}
        first = self.client.post("/api/v1/selection-results/generate", data=payload, content_type="application/json")
        second = self.client.post("/api/v1/selection-results/generate", data=payload, content_type="application/json")
        self.assertEqual(first.json()["code"], 0)
        self.assertEqual(second.json()["code"], 0)
        self.assertEqual(second.json()["data"]["count"], 2)

        # 唯一键 (event, player, stage, category) 使重复生成复用同一行：
        # 生效结果恒为 2 条，不会因重复生成产生重复行。
        active = SelectionResult.objects.filter(event=self.event, stage=2, category=2, status=1)
        inactive = SelectionResult.objects.filter(event=self.event, stage=2, category=2, status=0)
        self.assertEqual(active.count(), 2)
        self.assertEqual(inactive.count(), 0)
        self.assertEqual(active.order_by("rank")[0].rank, 1)
        self.assertEqual(active.order_by("rank")[0].player_id, self.players[0].id)

    def test_generate_does_not_touch_other_category(self) -> None:
        SelectionResult.objects.create(
            event=self.event_university, player=self.players[0], team_id=self.players[0].team_id,
            category=1, stage=2, total_score=Decimal("80.00"), rank=1, result=1, status=1,
        )
        self.client.post(
            "/api/v1/selection-results/generate",
            data={"event_id": self.event.id, "stage": 2, "category": 2, "top_n": 2},
            content_type="application/json",
        )
        university = SelectionResult.objects.filter(event=self.event_university, category=1, status=1)
        self.assertEqual(university.count(), 1)

    def test_generate_without_scores_returns_40001(self) -> None:
        body = self.client.post(
            "/api/v1/selection-results/generate",
            data={"event_id": self.event_empty.id, "stage": 1, "category": 3},
            content_type="application/json",
        ).json()
        self.assertEqual(body["code"], 40001)

    def test_results_list_after_generate(self) -> None:
        self.client.post(
            "/api/v1/selection-results/generate",
            data={"event_id": self.event.id, "stage": 2, "category": 2, "top_n": 3},
            content_type="application/json",
        )
        body = self.client.get(
            f"/api/v1/selection-results?event_id={self.event.id}&stage=2&category=2"
        ).json()
        self.assertEqual(body["code"], 0)
        self.assertEqual(body["data"]["total"], 3)
        self.assertEqual(body["data"]["list"][0]["rank"], 1)
        self.assertEqual(body["data"]["list"][0]["generated_by_name"], "系统管理员")



